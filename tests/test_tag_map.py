"""Единая карта тегов: ничего не потеряно и ничего не додумано."""

import pandas as pd
import pytest

from neftekod import config, parse_lims_pak, tag_map


def test_no_record_lost_from_source_sheets():
    """Записей в карте ровно столько, сколько непустых записей на листах.

    Пересчёт по листам сделан отдельной функцией «в лоб» — если сборщик
    пропустит блок (например, второй справочник на листе «КИП»), числа
    разойдутся.
    """
    frame = tag_map.build_tag_map()
    counts = tag_map.sheet_record_counts()
    assert counts == {"КИП": 97, "ЛА": 56, "ПАК": 2, "ВАК": 17}
    assert len(frame) == sum(counts.values())
    assert frame["source_sheet"].value_counts().to_dict() == counts


def test_kip_codes_are_qualified_by_installation():
    """`T6` есть и на АВТ, и на 24-2000 — короткий код не может быть ключом."""
    frame = tag_map.build_tag_map()
    kip = frame[frame["source_sheet"] == "КИП"]
    assert kip["code"].is_unique
    assert "АВТ:T6" in set(kip["code"])
    assert "24-2000:T6" in set(kip["code"])


def test_tags_without_unit_are_flagged_not_dropped():
    """Тег без подтверждённой единицы остаётся в карте с пустым `unit`."""
    frame = tag_map.attach_units_from_exports(
        tag_map.build_tag_map(), parse_lims_pak.parse_pak()
    )
    without_unit = tag_map.tags_without_unit(frame)
    # Единицы в пакете объявлены только в самих выгрузках ЛИМС/ПАК; из
    # справочника единицу получает только `24-2000:D15` (у тега серы имя в
    # справочнике другое — ADR-001, отождествление не делается).
    assert set(frame[frame["unit"].notna()]["code"]) == {"24-2000:D15"}
    assert len(without_unit) == len(frame) - 1
    assert set(without_unit["source_sheet"]) == {"КИП", "ЛА", "ПАК", "ВАК"}


def test_blocked_formulas_carry_adr_reference():
    """Формулы ВАК помечены ADR-002/003, а не выглядят готовыми к применению."""
    frame = tag_map.build_tag_map()
    vak = frame[frame["source_sheet"] == "ВАК"]
    assert vak["notes"].str.contains("ADR-00").all()
    pipeline = vak[vak["description"].str.contains("Pipeline")]
    assert set(pipeline["code"]) == {"24-2000:GODT:D15", "24-2000:GODT:T95"}
    assert pipeline["notes"].str.contains("ADR-003").all()


def test_duplicate_descriptions_are_reported():
    """Коды с одинаковым описанием находятся, а не теряются в справочнике."""
    duplicates = tag_map.duplicate_descriptions(tag_map.build_tag_map())
    found = {
        (row.installation, tuple(row.codes)) for row in duplicates.itertuples()
    }
    assert found == {
        ("24-2000", ("F1", "T12")),
        ("24-2000", ("F17", "T23")),
        ("АВТ", ("F56", "F57")),
        ("АВТ", ("F68", "F69")),
    }


def test_duplicate_description_pairs_disagree_in_telemetry():
    """Историческое наблюдение ADR-007: одинаковые описания, разные ряды.

    Это не доказательство неверного соответствия тегов: организаторы его
    подтвердили. Единицы и масштабы остаются неуточнёнными.
    """
    for path in (config.HT_TELEMETRY_CSV, config.AVT_TELEMETRY_CSV):
        if not path.exists():
            pytest.skip(f"нет {path} — скачать по ссылке из README")
    telemetry = {
        "24-2000": pd.read_csv(config.HT_TELEMETRY_CSV),
        "АВТ": pd.read_csv(config.AVT_TELEMETRY_CSV),
    }
    duplicates = tag_map.duplicate_descriptions(tag_map.build_tag_map())
    correlations = {}
    for row in duplicates.itertuples():
        frame = telemetry[row.installation]
        first, second = row.codes
        correlations[(row.installation, first, second)] = frame[first].corr(frame[second])

    agreeing = {key for key, value in correlations.items() if value > 0.95}
    assert agreeing == {("АВТ", "F56", "F57")}
    # Эти наблюдения сохранены для диагностики масштаба, без вывода о
    # неверном сопоставлении тегов (обновлённый ADR-007).
    for key in set(correlations) - agreeing:
        assert correlations[key] < 0.7


def test_la_sheet_has_no_machine_codes():
    """Лист «ЛА» — русские наименования анализов; кода показателя там нет."""
    frame = tag_map.build_tag_map()
    la = frame[frame["source_sheet"] == "ЛА"]
    assert la["code"].isna().all()
    assert la["sample_point"].notna().all()
    assert pd.notna(la["description"]).all()


@pytest.mark.parametrize("code, inputs, expected", [
    ("24-2000:GODT:T90", {
        "T12": 175.09552001953125, "F15": 2787.378662109375,
        "W7": 0.1431074738502502, "T23": 234.1968536376953,
        "F1": 2.613270998001098, "F26": 201.23069763183597,
    }, 324.91889751360736),
    ("24-2000:GODT:T50", {
        "P13": 3.758676290512085, "F9": 171.0937957763672, "T6": 360.1314697265625,
    }, 263.86193737838744),
    ("24-2000:GODT:CloudPoint", {
        "F22": 10507.5546875, "W7": 0.1431074738502502, "F25": 13242.359375,
        "F1": 2.613270998001098, "T6": 360.1314697265625,
        "F9": 171.0937957763672, "T16": 30.01093101501465,
    }, -0.9103307541167678),
    ("24-2000:GODT:CFPP", {
        "T23": 234.1968536376953, "T6": 360.1314697265625,
        "P8": 0.1171888560056686, "F9": 171.0937957763672,
        "W7": 0.1431074738502502, "P24": 0.5948644280433655,
    }, -17.328858721036617),
    ("24-2000:GODT:T95", {
        "F9": 100, "F2": 1000, "T6": 300, "lab_t95": 350,
    }, 313.7165),
    ("AVT6:240-350:CFPP", {
        "T33": 100, "P67": 1, "P4": 2, "F65": 12, "F32": 3, "F30": 5,
    }, 21.54194),
])
def test_exported_vak_corrections_produce_control_values(code, inputs, expected):
    """Проверка подстановки ловит старые коэффициенты, тег T6 и неверные скобки."""
    row = tag_map.build_tag_map().set_index("code").loc[code]
    formula = row["description"].replace("LIMS:24-2000.Pipeline.95%.T", "lab_t95")
    # Вычисляем только шесть известных выражений; это проверка экспорта,
    # не универсальный исполнитель произвольных формул для приложения.
    formula = formula.replace(",", ".").replace("×", "*").replace("x", "*")
    # Старая формула B5 синтаксически испорчена. Сначала требуем валидную
    # арифметическую запись, чтобы падение описывало исходную проблему.
    import ast
    try:
        expression = ast.parse(formula, mode="eval")
    except SyntaxError:
        pytest.fail(f"Невалидная формула {code}: {formula}")
    assert all(isinstance(node, (
        ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Name, ast.Load,
        ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd,
    )) for node in ast.walk(expression))
    actual = eval(compile(expression, "<vak-test>", "eval"), {"__builtins__": {}}, inputs)
    assert actual == pytest.approx(expected, abs=1e-8)


def test_formula_corrections_preserve_source_and_unresolved_pipeline():
    from openpyxl import load_workbook

    frame = tag_map.build_tag_map()
    assert {"formula_original", "formula_source", "formula_status"} <= set(frame)
    vak = frame[frame["source_sheet"] == "ВАК"].set_index("code")
    workbook = load_workbook(config.TAGS_XLSX, read_only=True, data_only=True)
    try:
        sheet = workbook["ВАК"]
        for cells in sheet.iter_rows(min_row=2, values_only=True):
            for col in (0, 2, 4):
                code = cells[col]
                if code:
                    assert vak.loc[code, "formula_original"] == " ".join(cells[col + 1].split())
    finally:
        workbook.close()
    corrected = vak[vak["formula_source"] == "organizer_clarifications_2026-09-10"]
    assert set(corrected.index) == {
        "24-2000:GODT:T90", "24-2000:GODT:T50", "24-2000:GODT:CloudPoint",
        "24-2000:GODT:CFPP", "24-2000:GODT:T95", "AVT6:240-350:CFPP",
    }
    assert corrected.loc["24-2000:GODT:T90", "formula_status"] == "corrected_unvalidated"
    blocked = vak[vak["formula_status"] == "blocked_pipeline"]
    assert set(blocked.index) == {"24-2000:GODT:D15", "24-2000:GODT:T95"}
    unchanged = vak[vak["formula_source"] == "source_xlsx"]
    assert unchanged["description"].equals(unchanged["formula_original"])
