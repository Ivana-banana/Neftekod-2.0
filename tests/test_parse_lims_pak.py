"""Тесты парсера на реальных файлах ЛИМС/ПАК, не на синтетике.

Синтетический пример не поймал бы главную ловушку формата: у каждого
показателя своя пара колонок с независимым охватом, и ошибка выравнивания
даёт правдоподобные, но чужие значения.
"""

import pandas as pd
import pytest

from neftekod import parse_lims_pak

# Заголовок LIMS заявляет «Количество значений:» — для двух показателей
# заявленное число на 1 больше, чем строк в самом файле. Проверено вручную
# по исходным ячейкам: лишней строки в файле нет, расхождение внесено
# экспортом. Фиксируем как факт данных, чтобы новое расхождение не утонуло.
KNOWN_COUNT_MISMATCH = {
    "LIMS:АВТ.1.CloudPoint": (14, 13),
    "LIMS:АВТ.1.CloudPoint_1": (113, 112),
}


def test_lims_series_are_split_by_indicator(lims_long):
    """54 показателя = 108 колонок исходника / 2 (пара «дата, значение»)."""
    assert lims_long["tag"].nunique() == 54
    assert set(lims_long["source"]) == {"lims"}
    assert set(lims_long["sample_point"]) == {
        "АВТ.1",
        "АВТ.2",
        "АВТ.2.1",
        "АВТ.3",
        "Гидроочистка.1",
        "Гидроочистка.2",
    }


def test_lims_counts_match_declared_header(lims_long):
    """Разобранное число точек сверяется с «Количество значений:» из файла."""
    declared = parse_lims_pak.lims_declared_counts()
    parsed = lims_long.groupby("tag").size().to_dict()
    assert set(declared) == set(parsed)

    mismatched = {
        tag: (declared[tag], parsed[tag]) for tag in declared if declared[tag] != parsed[tag]
    }
    assert mismatched == KNOWN_COUNT_MISMATCH


def test_lims_values_are_numeric_or_explicit_nan(lims_long):
    """'Pt Created' и подобные маркеры становятся NaN, а не 0 и не строкой."""
    assert pd.api.types.is_float_dtype(lims_long["value"])
    # Их мало и они локализованы — если доля вырастет, парсер молчать не должен.
    assert lims_long["value"].isna().sum() == 6


def test_pak_series_keep_independent_coverage(pak_long):
    """Ряды ПАК не выровнены по строкам: у D15 охват начинается на 2 года позже."""
    first_seen = pak_long.groupby("tag")["first_seen"].first()
    assert first_seen["24-2000:Mg.Sulfur"] == pd.Timestamp("2023-01-01 00:00:00")
    assert first_seen["24-2000:D15"] == pd.Timestamp("2025-03-05 10:20:00")

    d15 = pak_long[pak_long["tag"] == "24-2000:D15"]
    # Если бы пары колонок читались как одна таблица, у D15 появились бы
    # значения с датами серы (2023) — то есть чужой ряд под своим именем.
    assert d15["timestamp"].min() == first_seen["24-2000:D15"]
    assert d15["value"].between(700, 1000).all()


def test_pak_sulfur_tag_is_not_silently_renamed(pak_long):
    """ADR-001: имя тега берётся из файла, отождествление с `.Q` не делается."""
    tags = set(pak_long["tag"])
    assert parse_lims_pak.PAK_TAG_IN_FILE in tags
    assert parse_lims_pak.PAK_TAG_IN_REFERENCE_ONLY not in tags


def test_first_seen_is_per_tag_not_global(lims_long, pak_long):
    """`first_seen` — охват конкретного показателя, а не общая дата выгрузки."""
    frame = pd.concat([lims_long, pak_long], ignore_index=True)
    per_tag = frame.groupby("tag")["first_seen"].nunique()
    assert (per_tag == 1).all()
    assert frame["first_seen"].nunique() > 1
    assert (frame.groupby("tag")["timestamp"].min() == frame.groupby("tag")["first_seen"].first()).all()


def test_original_unit_conflicts_remain_visible(lims_long):
    """Единицы блока «АВТ.1» не совпадают с остальными блоками — см. ADR-005.

    Исправление нормализованной единицы не должно затирать исходный заголовок.
    """
    conflicts = parse_lims_pak.unit_conflicts(lims_long)
    assert set(conflicts["indicator"]) == {"50%.T", "D15", "EBP.T", "I350"}
    avt1 = conflicts[conflicts["tag"].str.startswith("LIMS:АВТ.1.")]
    assert dict(zip(avt1["indicator"], avt1["unit_declared"])) == {
        "50%.T": "кг/м3",
        "D15": "°С",
        "EBP.T": "% об.",
        "I350": "°С",
    }


def test_sulfur_unit_difference_is_surfaced(lims_long, pak_long):
    """Сера объявлена как «ppm» в ПАК и «мг/кг» в ЛИМС — см. ADR-005.

    Величина одна и та же (1 ppm = 1 мг/кг для массовой доли), но это
    конвертация, а не совпадение имён: она касается жёсткой проверки
    «сера ≤ 10 мг/кг» и не должна применяться молча.
    """
    frame = pd.concat([lims_long, pak_long], ignore_index=True)
    conflicts = parse_lims_pak.unit_conflicts(frame)
    sulfur = conflicts[conflicts["indicator"] == "Mg.Sulfur"]
    assert dict(zip(sulfur["tag"], sulfur["unit_declared"])) == {
        "24-2000:Mg.Sulfur": "ppm",
        "LIMS:Гидроочистка.2.Mg.Sulfur": "мг/кг",
    }


@pytest.mark.parametrize("column", parse_lims_pak.LONG_COLUMNS)
def test_long_format_columns(lims_long, column):
    assert column in lims_long.columns


def test_lims_units_follow_indicators_without_changing_values(lims_long):
    """Сдвиг заголовка не должен менять число или давать плотность в °C."""
    from openpyxl import load_workbook
    from neftekod import config

    assert {"unit", "unit_basis"} <= set(lims_long)
    expected = {
        "50%.T": ("E5", "кг/м3", "°C"),
        "EBP.T": ("G5", "% об.", "°C"),
        "D15": ("K5", "°С", "кг/м3"),
        "I350": ("M5", "°С", "% об."),
    }
    workbook = load_workbook(config.LIMS_XLSX, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        for indicator, (date_cell, declared, unit) in expected.items():
            raw_date = sheet[date_cell]
            raw_value = sheet.cell(raw_date.row, raw_date.column + 1).value
            rows = lims_long[lims_long["tag"] == f"LIMS:АВТ.1.{indicator}"]
            assert set(rows["unit_declared"]) == {declared}
            assert set(rows["unit"]) == {unit}
            assert set(rows["unit_basis"]) == {"organizer_indicator"}
            point = rows[rows["timestamp"] == pd.Timestamp(raw_date.value)].iloc[0]
            assert point["value"] == pytest.approx(float(str(raw_value).replace(",", ".")))
    finally:
        workbook.close()


def test_mass_ppm_is_normalized_without_rescaling(pak_long, lims_long):
    from openpyxl import load_workbook
    from neftekod import config

    assert {"unit", "unit_basis"} <= set(pak_long)
    sulfur = pak_long[pak_long["tag"] == "24-2000:Mg.Sulfur"]
    lab = lims_long[lims_long["tag"] == "LIMS:Гидроочистка.2.Mg.Sulfur"]
    assert set(sulfur["unit"]) == set(lab["unit"]) == {"мг/кг"}
    assert set(sulfur["unit_declared"]) == {"ppm"}
    assert set(sulfur["unit_basis"]) == {"mass_ppm_equivalence"}
    workbook = load_workbook(config.PAK_XLSX, read_only=True, data_only=True)
    try:
        rows = list(workbook.worksheets[0].iter_rows(max_row=3, values_only=True))
        col = rows[0].index("24-2000:Mg.Sulfur")
        timestamp, raw_value = rows[2][col:col + 2]
        point = sulfur[sulfur["timestamp"] == pd.Timestamp(timestamp)].iloc[0]
        assert point["value"] == float(str(raw_value).replace(",", "."))
    finally:
        workbook.close()


def test_lims_availability_keeps_sample_time(lims_long):
    assert {"available_at", "availability_basis"} <= set(lims_long)
    assert (lims_long["available_at"] - lims_long["timestamp"]).eq(pd.Timedelta(hours=4)).all()
    assert set(lims_long["availability_basis"]) == {"sample_time_plus_4h_conservative"}
    # Охват по отбору не сдвигается вместе с доступностью признаков.
    sulfur = lims_long[lims_long["tag"] == "LIMS:Гидроочистка.2.Mg.Sulfur"]
    assert sulfur["timestamp"].min() == pd.Timestamp("2023-01-01 10:00")


def test_pak_availability_has_no_invented_lab_delay(pak_long):
    assert {"available_at", "availability_basis"} <= set(pak_long)
    assert pak_long["available_at"].equals(pak_long["timestamp"])
    assert set(pak_long["availability_basis"]) == {"measurement_time_assumed"}


def test_csv_export_roundtrips_all_times(lims_long, pak_long, tmp_path):
    """Формат даты не должен зависеть от соседних колонок или чанка pandas."""
    frame = pd.concat([lims_long, pak_long], ignore_index=True).sort_values(
        ["source", "tag", "timestamp"]
    ).reset_index(drop=True)
    path = tmp_path / "long.csv"
    parse_lims_pak.write_long_format(frame, path)
    times = ["timestamp", "available_at", "first_seen"]
    loaded = pd.read_csv(path, parse_dates=times)
    for column in times:
        assert pd.api.types.is_datetime64_any_dtype(loaded[column]), column
        pd.testing.assert_series_equal(loaded[column], frame[column], check_dtype=False)
