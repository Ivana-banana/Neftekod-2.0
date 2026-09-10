"""Парсер lims_export_2023.xlsx / pak_export_2023.xlsx в единый long-format.

Формат исходников описан в `.claude/skills/parse-lims-pak/SKILL.md`: это не
таблица для `pd.read_excel()`. У каждого показателя своя пара колонок
(дата, значение) с независимым охватом и частотой; строки разных показателей
между собой не выровнены.

Заголовки у двух файлов разные:

LIMS (`lims_export_2023.xlsx`)      PAK (`pak_export_2023.xlsx`)
  строка 1: установка + точка         строка 1: код тега
  строка 2: показатель                строка 2: единица
  строка 3: единица                   строка 3+: данные (дата, значение)
  строка 4: «Количество значений:» N
  строка 5+: данные (дата, значение)

Результат — long-format `(timestamp, tag, value, source, ...)`, где на каждое
значение сохранён `first_seen` — первая дата наличия этого показателя.
`timestamp` остаётся временем отбора/измерения, `available_at` — оценка
доступности для признаков. `unit_declared` сохраняет исходную единицу,
`unit` — нормализованную; численные значения не масштабируются.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from . import config

log = logging.getLogger(__name__)

LONG_COLUMNS = [
    "timestamp",
    "tag",
    "value",
    "source",
    "unit_declared",
    "sample_point",
    "indicator",
    "first_seen",
    "unit",
    "unit_basis",
    "available_at",
    "availability_basis",
]

# Организаторы подтвердили ошибку строки единиц: названия показателей верны.
# См. ADR-005. Это исправление метаданных, не масштабирование значений.
LIMS_INDICATOR_UNITS = {
    "50%.T": "°C",
    "EBP.T": "°C",
    "D15": "кг/м3",
    "I350": "% об.",
}

# Справочник «ПАК» (docs/tz/tags_hackathon.md) называет тег серы `24-2000:Mg.Sulfur.Q`,
# в самом pak_export_2023.xlsx колонка называется `24-2000:Mg.Sulfur`.
# Тождество не подтверждено организаторами — см. ADR-001. Молча переименовывать
# нельзя: тег напрямую участвует в жёсткой проверке «сера ≤ 10 мг/кг».
PAK_TAG_IN_REFERENCE_ONLY = "24-2000:Mg.Sulfur.Q"
PAK_TAG_IN_FILE = "24-2000:Mg.Sulfur"


@dataclass(frozen=True)
class SeriesHeader:
    """Описание одной пары колонок (дата, значение) в исходном файле."""

    tag: str
    indicator: str
    sample_point: str
    unit_declared: str | None
    col_date: int  # 0-based индекс колонки с датой
    declared_count: int | None  # «Количество значений:» — только у LIMS


def _clean(value) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _to_float(value) -> float:
    """Число из ячейки. Экспорт хранит часть значений строками, часть — числами."""
    if value is None:
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not text:
        return float("nan")
    try:
        return float(text)
    except ValueError:
        return float("nan")


def _to_timestamp(value):
    if isinstance(value, datetime):
        return value
    if value is None:
        return None
    try:
        return pd.to_datetime(value)
    except (ValueError, TypeError):
        return None


def _parse_sample_point(raw: str) -> str:
    """`Установка 'Гидроочистка'.. Точка отбора '2'. Продукт '…'` → `Гидроочистка.2`.

    Точка/установка берутся как есть из файла: соответствие обозначениям
    организаторов (`LIMS:24-2000.Pipeline.*` из формул ВАК) не установлено —
    см. ADR-003, поэтому переименование в их нотацию здесь не делается.
    """
    parts = [p.strip(" .") for p in raw.split("'")]
    installation = parts[1] if len(parts) > 1 else raw
    point = parts[3] if len(parts) > 3 else "?"
    return f"{installation}.{point}"


def _read_lims_headers(rows: list[tuple]) -> list[SeriesHeader]:
    header_point, header_indicator, header_unit, header_count = rows[:4]
    current_point = None
    headers: list[SeriesHeader] = []
    for col in range(len(header_indicator)):
        point_cell = _clean(header_point[col])
        if point_cell:
            current_point = _parse_sample_point(point_cell)
        indicator = _clean(header_indicator[col])
        if not indicator or current_point is None:
            continue
        declared = _clean(header_count[col + 1]) if col + 1 < len(header_count) else None
        headers.append(
            SeriesHeader(
                tag=f"LIMS:{current_point}.{indicator}",
                indicator=indicator,
                sample_point=current_point,
                unit_declared=_clean(header_unit[col]),
                col_date=col,
                declared_count=int(float(declared)) if declared else None,
            )
        )
    return headers


def _read_pak_headers(rows: list[tuple]) -> list[SeriesHeader]:
    header_tag, header_unit = rows[:2]
    headers: list[SeriesHeader] = []
    for col in range(len(header_tag)):
        tag = _clean(header_tag[col])
        if not tag:
            continue
        headers.append(
            SeriesHeader(
                tag=tag,
                indicator=tag.split(":")[-1],
                sample_point=tag.split(":")[0],
                unit_declared=_clean(header_unit[col]),
                col_date=col,
                declared_count=None,
            )
        )
    return headers


def _extract(path: Path, source: str, header_rows: int, headers_fn) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        head = [row for _, row in zip(range(header_rows), sheet.iter_rows(values_only=True))]
        headers = headers_fn(head)
        buckets: dict[str, list[tuple]] = {h.tag: [] for h in headers}
        if len(buckets) != len(headers):
            raise ValueError(f"{path.name}: повторяющиеся коды тегов в заголовке")

        for row in sheet.iter_rows(min_row=header_rows + 1, values_only=True):
            for header in headers:
                cell_date = row[header.col_date]
                if cell_date is None:
                    continue
                timestamp = _to_timestamp(cell_date)
                if timestamp is None:
                    continue
                buckets[header.tag].append((timestamp, _to_float(row[header.col_date + 1])))
    finally:
        workbook.close()

    frames = []
    for header in headers:
        points = buckets[header.tag]
        if not points:
            log.warning("%s: показатель %s без значений", source, header.tag)
            continue
        frame = pd.DataFrame(points, columns=["timestamp", "value"])
        frame["tag"] = header.tag
        frame["source"] = source
        frame["unit_declared"] = header.unit_declared
        frame["sample_point"] = header.sample_point
        frame["indicator"] = header.indicator
        frame["unit"] = {"°С": "°C"}.get(header.unit_declared, header.unit_declared)
        frame["unit_basis"] = "declared"
        if source == "lims":
            if header.indicator in LIMS_INDICATOR_UNITS:
                frame["unit"] = LIMS_INDICATOR_UNITS[header.indicator]
                frame["unit_basis"] = "organizer_indicator"
            # См. ADR-006: не переносим время цели на время её публикации.
            frame["available_at"] = frame["timestamp"] + pd.Timedelta(
                hours=config.LIMS_PUBLICATION_DELAY_HOURS
            )
            frame["availability_basis"] = "sample_time_plus_4h_conservative"
        else:
            # Для ПАК задержка не предоставлена; это явное допущение,
            # а не подтверждённый момент публикации.
            frame["available_at"] = frame["timestamp"]
            frame["availability_basis"] = "measurement_time_assumed"
            if header.tag == PAK_TAG_IN_FILE and header.unit_declared == "ppm":
                # См. ADR-001/005: только подтверждённая массовая сера,
                # 1 ppm = 1 мг/кг. Не распространять на неизвестные ppm.
                frame["unit"] = "мг/кг"
                frame["unit_basis"] = "mass_ppm_equivalence"
        # Первая дата наличия показателя: охват у каждого свой (например,
        # 24-2000:D15 начинается только с 2025-03-05), и через такой разрыв
        # интерполировать нельзя — см. SKILL.md.
        frame["first_seen"] = frame["timestamp"].min()
        non_numeric = int(frame["value"].isna().sum())
        if non_numeric:
            # Экспорт кладёт в колонку значения служебные маркеры PI
            # (встречается 'Pt Created'). Это не число и не ноль — NaN,
            # с явным сообщением, а не тихая подмена.
            log.warning("%s: %s — %d нечисловых значений", source, header.tag, non_numeric)
        if header.declared_count is not None and header.declared_count != len(frame):
            log.warning(
                "%s: %s — в заголовке заявлено %d значений, разобрано %d",
                source,
                header.tag,
                header.declared_count,
                len(frame),
            )
        frames.append(frame[LONG_COLUMNS])

    return pd.concat(frames, ignore_index=True)


def parse_lims(path: Path | None = None) -> pd.DataFrame:
    return _extract(path or config.LIMS_XLSX, "lims", 4, _read_lims_headers)


def parse_pak(path: Path | None = None) -> pd.DataFrame:
    frame = _extract(path or config.PAK_XLSX, "pak", 2, _read_pak_headers)
    tags = set(frame["tag"])
    if PAK_TAG_IN_FILE in tags and PAK_TAG_IN_REFERENCE_ONLY not in tags:
        # см. ADR-001 — расхождение имени со справочником логируется явно,
        # переименование не выполняется до ответа организаторов.
        log.warning(
            "ADR-001: в файле есть %s, в справочнике «ПАК» — %s; "
            "теги не отождествляются, используется имя из файла",
            PAK_TAG_IN_FILE,
            PAK_TAG_IN_REFERENCE_ONLY,
        )
    return frame


def declared_counts(path: Path, header_rows: int, headers_fn) -> dict[str, int | None]:
    """Строка «Количество значений:» из заголовка — независимая сверка парсера."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        head = [row for _, row in zip(range(header_rows), sheet.iter_rows(values_only=True))]
    finally:
        workbook.close()
    return {h.tag: h.declared_count for h in headers_fn(head)}


def lims_declared_counts(path: Path | None = None) -> dict[str, int | None]:
    return declared_counts(path or config.LIMS_XLSX, 4, _read_lims_headers)


def unit_conflicts(frame: pd.DataFrame) -> pd.DataFrame:
    """Показатели, у которых в разных блоках заявлены разные единицы.

    Исходные расхождения сохраняются для аудита после нормализации единиц
    по ответу организаторов (см. ADR-005).
    """
    per_indicator = (
        frame.drop_duplicates(["indicator", "unit_declared", "tag"])
        .groupby("indicator")["unit_declared"]
        .nunique(dropna=False)
    )
    conflicting = per_indicator[per_indicator > 1].index
    return (
        frame[frame["indicator"].isin(conflicting)]
        .drop_duplicates(["indicator", "unit_declared", "tag"])
        .loc[:, ["indicator", "unit_declared", "tag"] + [
            column for column in ("unit", "unit_basis") if column in frame
        ]]
        .sort_values(["indicator", "unit_declared", "tag"])
        .reset_index(drop=True)
    )


def build_long_format() -> pd.DataFrame:
    frame = pd.concat([parse_lims(), parse_pak()], ignore_index=True)
    return frame.sort_values(["source", "tag", "timestamp"]).reset_index(drop=True)


def write_long_format(frame: pd.DataFrame, path: Path) -> None:
    """Единый формат времени CSV, включая полночь и доли секунды.

    Автоформат pandas зависит от содержимого чанка и может смешать дату
    без времени с полной датой внутри first_seen, ломая read_csv(parse_dates).
    """
    frame.to_csv(path, index=False, date_format="%Y-%m-%d %H:%M:%S.%f")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    long_format = build_long_format()
    config.BUILD_DIR.mkdir(exist_ok=True)
    write_long_format(long_format, config.LONG_FORMAT_CSV)
    print(f"{len(long_format)} строк → {config.LONG_FORMAT_CSV}")
