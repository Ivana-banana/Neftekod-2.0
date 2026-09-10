"""Единая карта тегов: листы КИП, ЛА, ПАК, ВАК из docs/tz/tags_hackathon.xlsx в один файл.

У четырёх листов четыре разных формата (см. docs/tz/tags_hackathon.md), поэтому в
исходном виде справочник не годится ни для проверки колонок телеметрии, ни
для подстановки единиц. Здесь они сводятся к одной таблице:

    code | description | unit | source_sheet | installation | sample_point | notes

Единицы измерения есть не у всех источников. Тег без подтверждённой единицы
получает `unit = None` и попадает в `tags_without_unit()` — молча
«догадываться» о физическом смысле по букве кода запрещено (ТЗ, §2).
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from . import config
from .formula_corrections import FORMULA_SOURCE, VAK_CORRECTIONS

MAP_COLUMNS = [
    "code",
    "description",
    "unit",
    "source_sheet",
    "installation",
    "sample_point",
    "notes",
    "formula_original",
    "formula_source",
    "formula_status",
]


def _clean(value) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _sheet_rows(path: Path) -> dict[str, list[tuple]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {ws.title: list(ws.iter_rows(values_only=True)) for ws in workbook.worksheets}
    finally:
        workbook.close()


def _kip_records(rows: list[tuple]) -> list[dict]:
    """Лист «КИП»: два независимых справочника рядом — АВТ и 24-2000.

    Коды (`T1`, `F1`, …) уникальны только внутри установки, поэтому
    ключом служит `installation:code`.
    """
    header = rows[0]
    blocks = [(0, 1, _clean(header[1])), (2, 3, _clean(header[3]))]
    records = []
    for desc_col, code_col, installation in blocks:
        for row in rows[1:]:
            code = _clean(row[code_col])
            description = _clean(row[desc_col])
            if not code:
                continue
            records.append(
                {
                    "code": f"{installation}:{code}",
                    "description": description,
                    "unit": None,
                    "source_sheet": "КИП",
                    "installation": installation,
                    "sample_point": None,
                    "notes": f"короткий код в телеметрии: {code}",
                }
            )
    return records


def _la_records(rows: list[tuple]) -> list[dict]:
    """Лист «ЛА»: колонка = точка отбора, ячейка = наименование анализа.

    Машиночитаемого кода на этом листе нет — только русские названия. Связь
    с кодами показателей из lims_export_2023.xlsx (`50%.T`, `CFPP`, …)
    справочником не задана, поэтому здесь не выдумывается.
    """
    header = rows[0]
    records = []
    for col, point_cell in enumerate(header):
        point = _clean(point_cell)
        if not point:
            continue
        for row in rows[1:]:
            name = _clean(row[col]) if col < len(row) else None
            if not name:
                continue
            records.append(
                {
                    "code": None,
                    "description": name,
                    "unit": None,
                    "source_sheet": "ЛА",
                    "installation": point.split("'")[1] if "'" in point else None,
                    "sample_point": point,
                    "notes": "лист «ЛА» не содержит кода показателя",
                }
            )
    return records


def _pak_records(rows: list[tuple]) -> list[dict]:
    installation = _clean(rows[0][0])
    records = []
    for row in rows[1:]:
        code = _clean(row[1])
        if not code:
            continue
        note = None
        if code == "24-2000:Mg.Sulfur.Q":
            # см. ADR-001
            note = "в pak_export_2023.xlsx колонка называется 24-2000:Mg.Sulfur (ADR-001)"
        records.append(
            {
                "code": code,
                "description": _clean(row[0]),
                "unit": None,
                "source_sheet": "ПАК",
                "installation": installation,
                "sample_point": None,
                "notes": note,
            }
        )
    return records


def _vak_records(rows: list[tuple]) -> list[dict]:
    """Лист «ВАК»: пары колонок (код анализатора, формула) под общим заголовком."""
    header = rows[0]
    records = []
    for col in range(0, len(header), 2):
        installation = _clean(header[col])
        if not installation:
            continue
        for row in rows[1:]:
            code = _clean(row[col])
            if not code:
                continue
            original = _clean(row[col + 1])
            formula = VAK_CORRECTIONS.get(code, original)
            corrected = code in VAK_CORRECTIONS
            status = "corrected_unvalidated" if corrected else "original_unvalidated"
            note = "формула не проверена: ADR-002"
            if corrected:
                note = "исправлено по ответу организаторов; точность не проверена: ADR-002"
            if formula and "Pipeline" in formula:
                status = "blocked_pipeline"
                note += "; точка Pipeline не сопоставлена с ЛИМС: ADR-003"
            records.append(
                {
                    "code": code,
                    "description": formula,
                    "unit": None,
                    "source_sheet": "ВАК",
                    "installation": installation,
                    "sample_point": None,
                    "notes": note,
                    "formula_original": original,
                    "formula_source": FORMULA_SOURCE if corrected else "source_xlsx",
                    "formula_status": status,
                }
            )
    return records


_BUILDERS = {
    "КИП": _kip_records,
    "ЛА": _la_records,
    "ПАК": _pak_records,
    "ВАК": _vak_records,
}


def build_tag_map(path: Path | None = None) -> pd.DataFrame:
    sheets = _sheet_rows(path or config.TAGS_XLSX)
    missing = set(_BUILDERS) - set(sheets)
    if missing:
        raise ValueError(f"в справочнике нет листов: {sorted(missing)}")
    records: list[dict] = []
    for sheet, builder in _BUILDERS.items():
        records.extend(builder(sheets[sheet]))
    return pd.DataFrame(records, columns=MAP_COLUMNS)


def sheet_record_counts(path: Path | None = None) -> dict[str, int]:
    """Независимый пересчёт записей по листам — для сверки с картой.

    Считается «в лоб», без разбора структуры листов: непустые ячейки в
    колонках-кодах (КИП/ПАК/ВАК) и непустые ячейки под заголовком (ЛА).
    """
    sheets = _sheet_rows(path or config.TAGS_XLSX)
    counts: dict[str, int] = {}

    counts["КИП"] = sum(
        1 for row in sheets["КИП"][1:] for col in (1, 3) if _clean(row[col])
    )
    counts["ЛА"] = sum(
        1
        for row in sheets["ЛА"][1:]
        for col, head in enumerate(sheets["ЛА"][0])
        if _clean(head) and col < len(row) and _clean(row[col])
    )
    counts["ПАК"] = sum(1 for row in sheets["ПАК"][1:] if _clean(row[1]))
    counts["ВАК"] = sum(
        1
        for row in sheets["ВАК"][1:]
        for col in range(0, len(sheets["ВАК"][0]), 2)
        if _clean(sheets["ВАК"][0][col]) and _clean(row[col])
    )
    return counts


def attach_units_from_exports(tag_map: pd.DataFrame, long_format: pd.DataFrame) -> pd.DataFrame:
    """Проставить нормализованные единицы выгрузки ПАК по точному имени тега.

    Единицы берутся из данных, а не из справочника: в справочнике их нет.
    Строка `unit` остаётся пустой для КИП/ЛА/ВАК — подтверждённого источника
    единиц для них в пакете нет.
    """
    declared = (
        long_format[long_format["source"] == "pak"]
        .drop_duplicates("tag")
        .set_index("tag")["unit"]
    )
    result = tag_map.copy()
    result["unit"] = result["unit"].astype("object")
    matched = result["code"].map(declared)
    result["unit"] = result["unit"].where(matched.isna(), matched)
    return result


def tags_without_unit(tag_map: pd.DataFrame) -> pd.DataFrame:
    return tag_map[tag_map["unit"].isna()].reset_index(drop=True)


def duplicate_descriptions(tag_map: pd.DataFrame) -> pd.DataFrame:
    """Коды листа «КИП», у которых внутри установки совпадает описание.

    Описание не определяет колонку однозначно. Дубликат сам по себе не
    доказывает ошибку сопоставления: соответствие подтверждено организаторами,
    единицы и масштабы остаются открытыми (см. ADR-007).

        installation | description | codes
    """
    kip = tag_map[tag_map["source_sheet"] == "КИП"]
    grouped = (
        kip.groupby(["installation", "description"])["code"]
        .apply(lambda codes: sorted(code.split(":", 1)[1] for code in codes))
        .reset_index(name="codes")
    )
    duplicates = grouped[grouped["codes"].str.len() > 1]
    return duplicates.sort_values(["installation", "description"]).reset_index(drop=True)


if __name__ == "__main__":
    frame = build_tag_map()
    config.BUILD_DIR.mkdir(exist_ok=True)
    frame.to_csv(config.TAG_MAP_CSV, index=False)
    print(f"{len(frame)} записей → {config.TAG_MAP_CSV}")
    print(frame["source_sheet"].value_counts().to_string())
